from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def load_rules() -> dict[str, dict]:
    workbook = load_workbook(ROOT / "data/internal_rules.xlsx", read_only=True)
    try:
        rows = workbook["rules"].iter_rows(values_only=True)
        headers = next(rows)
        result = {}
        for values in rows:
            if all(value is None for value in values):
                continue
            rule = dict(zip(headers, values))
            code = rule["dept_code"]
            if code in result:
                raise ValueError(f"Trùng quy ước cho {code}")
            if rule["kind"] not in {"margin_below", "cost_growth_above", "revenue_growth_below"}:
                raise ValueError("Loại rule không được hỗ trợ")
            rule["threshold"] = float(rule["threshold"])
            result[code] = rule
        return result
    finally:
        workbook.close()


def growth(current: float, previous: float) -> float | None:
    return None if previous == 0 else (current - previous) / previous * 100


def evaluate(current: list[dict], previous: list[dict]) -> list[dict]:
    rules = load_rules()
    prev_by_code = {row["dept_code"]: row for row in previous}
    result = []
    for row in current:
        rule = rules[row["dept_code"]]  # Thiếu rule phải fail, không đoán.
        prev = prev_by_code.get(row["dept_code"])
        revenue, cost = row["revenue_usd"], row["cost_usd"]
        margin = None if revenue == 0 else (revenue - cost) / revenue * 100
        cost_growth = None if prev is None else growth(cost, prev["cost_usd"])
        revenue_growth = None if prev is None else growth(revenue, prev["revenue_usd"])
        kind = rule["kind"]
        value = {
            "margin_below": margin,
            "cost_growth_above": cost_growth,
            "revenue_growth_below": revenue_growth,
        }[kind]
        if value is None:
            alert = "KHÔNG ĐỦ DỮ LIỆU"
        else:
            violated = value > rule["threshold"] if kind.endswith("above") else value < rule["threshold"]
            alert = rule["label"] if violated else "ĐẠT"
        result.append(
            {
                **row,
                "dept_name": rule["dept_name"],
                "rule_id": rule["rule_id"],
                "margin_pct": margin,
                "cost_growth_pct": cost_growth,
                "revenue_growth_pct": revenue_growth,
                "alert": alert,
            }
        )
    return result
