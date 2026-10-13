import pytest
from openpyxl import load_workbook

from excel_agent.core.exceptions import ValidationError
from excel_agent.mock_data_api.service import DataRepository
from excel_agent.schemas.data_api import ExportRequest, Filter
from excel_agent.schemas.workbook import ApplyRequest
from excel_agent.services.workbook_service import WorkbookService


@pytest.fixture
def hanoi(repo: DataRepository, workbooks: WorkbookService) -> str:
    content, name, _ = repo.export_xlsx(
        ExportRequest(dataset="khach_hang", filters=[Filter(column="Thành phố", value="Hà Nội")])
    )
    return workbooks.create(content, name, {"dataset": "khach_hang"}).id


def _col(workbooks: WorkbookService, wid: str, name: str) -> list:
    pv = workbooks.preview(wid)
    i = pv.columns.index(name)
    return [row[i] for row in pv.rows]


def test_preview_formats_dates_and_keeps_numbers(workbooks: WorkbookService, hanoi: str) -> None:
    pv = workbooks.preview(hanoi, limit=2)
    assert pv.version == 1 and len(pv.rows) == 2 and pv.total_rows == 9
    row = dict(zip(pv.columns, pv.rows[0], strict=True))
    assert isinstance(row["Hạn mức nợ (VND)"], int)
    assert row["Ngày tạo"].count("/") == 2


def test_set_column_one_value_many_rows(repo: DataRepository, workbooks: WorkbookService, hanoi: str) -> None:
    rec = repo.search("khach_hang", "Đà Nẵng", column="Thành phố").records[0]
    res = workbooks.apply(
        hanoi, ApplyRequest(mode="set_column", records=[rec], target_column="thanh pho", row_indexes=[0, 2])
    )
    assert res.changed_rows == [0, 2] and res.workbook.version == 2
    assert _col(workbooks, hanoi, "Thành phố")[:3] == ["Đà Nẵng", "Hà Nội", "Đà Nẵng"]


def test_set_column_pairwise_and_mismatch(repo: DataRepository, workbooks: WorkbookService, hanoi: str) -> None:
    recs = repo.search("khach_hang", "Cần Thơ", column="Thành phố").records[:2]
    workbooks.apply(hanoi, ApplyRequest(mode="set_column", records=recs, target_column="Email", row_indexes=[1, 3]))
    emails = _col(workbooks, hanoi, "Email")
    assert [emails[1], emails[3]] == [r["Email"] for r in recs]
    with pytest.raises(ValidationError, match="kết quả cho 3 dòng"):
        workbooks.apply(
            hanoi, ApplyRequest(mode="set_column", records=recs, target_column="Email", row_indexes=[0, 1, 2])
        )


def test_set_column_creates_new_column(repo: DataRepository, workbooks: WorkbookService, hanoi: str) -> None:
    rec = repo.search("khach_hang", "KH-0001").records[0]
    res = workbooks.apply(
        hanoi,
        ApplyRequest(
            mode="set_column",
            records=[rec],
            target_column="Người phụ trách",
            source_field="Tên khách hàng",
            row_indexes=[0],
            create_column=True,
        ),
    )
    assert "Người phụ trách" in res.workbook.columns and res.warnings


def test_validation_blocks_bad_writes(repo: DataRepository, workbooks: WorkbookService, hanoi: str) -> None:
    rec = repo.search("khach_hang", "KH-0001").records[0]
    with pytest.raises(ValidationError, match="không phải ngày"):
        workbooks.apply(
            hanoi,
            ApplyRequest(
                mode="set_column", records=[rec], target_column="Ngày tạo", source_field="Email", row_indexes=[0]
            ),
        )
    with pytest.raises(ValidationError, match="Chưa chọn dòng"):
        workbooks.apply(hanoi, ApplyRequest(mode="set_column", records=[rec], target_column="Email"))
    with pytest.raises(ValidationError, match="ngoài phạm vi"):
        workbooks.apply(hanoi, ApplyRequest(mode="set_column", records=[rec], target_column="Email", row_indexes=[99]))
    with pytest.raises(ValidationError, match="phiên bản"):
        workbooks.apply(
            hanoi,
            ApplyRequest(mode="set_column", records=[rec], target_column="Email", row_indexes=[0], expected_version=7),
        )


def test_append_rows_extends_table_and_blocks_duplicates(
    repo: DataRepository, workbooks: WorkbookService, hanoi: str
) -> None:
    recs = repo.search("khach_hang", "Cần Thơ", column="Thành phố").records[:2]
    res = workbooks.apply(hanoi, ApplyRequest(mode="append_rows", records=recs))
    assert res.changed_rows == [9, 10] and res.workbook.rows == 11
    ws = load_workbook(workbooks.file_path(hanoi)).worksheets[0]
    assert ws.max_row == 12 and list(ws.tables.values())[0].ref == "A1:H12"
    assert ws["A11"].value == recs[0]["Mã KH"] and ws["H11"].number_format == "dd/mm/yyyy"
    with pytest.raises(ValidationError, match="đã có trong bảng"):
        workbooks.apply(hanoi, ApplyRequest(mode="append_rows", records=recs[:1]))


def test_undo_returns_previous_version(repo: DataRepository, workbooks: WorkbookService, hanoi: str) -> None:
    with pytest.raises(ValidationError, match="hoàn tác"):
        workbooks.undo(hanoi)
    recs = repo.search("khach_hang", "Huế", column="Thành phố").records[:1]
    workbooks.apply(hanoi, ApplyRequest(mode="append_rows", records=recs))
    info = workbooks.undo(hanoi)
    assert info.version == 1 and info.rows == 9
