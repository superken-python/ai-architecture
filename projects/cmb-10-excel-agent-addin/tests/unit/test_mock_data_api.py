import io

import pytest
from openpyxl import load_workbook

from excel_agent.core.exceptions import NotFoundError
from excel_agent.mock_data_api.service import DataRepository
from excel_agent.schemas.data_api import ExportRequest, Filter


def test_search_by_column_and_full_text(repo: DataRepository) -> None:
    by_col = repo.search("khach_hang", "da nang", column="thanh pho")
    assert by_col.column == "Thành phố"
    assert by_col.total > 0 and all(r["Thành phố"] == "Đà Nẵng" for r in by_col.records)
    full = repo.search("Khách hàng", "KH-0003")
    assert full.records[0]["Mã KH"] == "KH-0003"


def test_search_unknown_column_or_dataset(repo: DataRepository) -> None:
    with pytest.raises(NotFoundError):
        repo.search("khach_hang", "x", column="Không có")
    with pytest.raises(NotFoundError):
        repo.search("khong_co", "x")


def test_filters_numeric_and_contains(repo: DataRepository) -> None:
    rows = repo.query("san_pham", [Filter(column="Đơn giá (VND)", op="gte", value="10.000.000")])
    assert rows and all(r["Đơn giá (VND)"] >= 10_000_000 for r in rows)
    rows = repo.query("san_pham", [Filter(column="Tên sản phẩm", op="contains", value="laptop")])
    assert {r["Mã SP"] for r in rows} == {"SP-001", "SP-002"}


def test_export_creates_formatted_table(repo: DataRepository) -> None:
    content, name, n = repo.export_xlsx(
        ExportRequest(dataset="khach_hang", filters=[Filter(column="Thành phố", value="Hà Nội")])
    )
    assert name.startswith("khach_hang_") and name.endswith(".xlsx") and n > 0
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Khách hàng"]
    assert ws.max_row == n + 1 and ws["A1"].value == "Mã KH"
    assert list(ws.tables.values())[0].ref == f"A1:H{n + 1}"
    assert "Thông tin" in wb.sheetnames


def test_export_without_rows_has_no_table(repo: DataRepository) -> None:
    content, _, n = repo.export_xlsx(
        ExportRequest(dataset="khach_hang", filters=[Filter(column="Mã KH", value="KHONG")])
    )
    ws = load_workbook(io.BytesIO(content)).worksheets[0]
    assert n == 0 and not ws.tables and ws.max_row == 1
