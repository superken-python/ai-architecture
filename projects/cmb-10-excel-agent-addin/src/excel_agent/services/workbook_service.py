"""Quản lý workbook: lưu theo phiên bản, xem dữ liệu, áp kết quả tìm kiếm, hoàn tác.

Mọi thao tác ghi đi qua lớp kiểm chứng xác định (EFF-07) và tạo một phiên bản mới, nên hoàn tác
chỉ là quay về phiên bản trước. Trong add-in Excel thật, phần này chạy bằng Office.js (xem docs/02).
"""

from __future__ import annotations

import json
import threading
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from excel_agent.core.exceptions import NotFoundError, ValidationError
from excel_agent.core.text import find_name, normalize
from excel_agent.core.values import parse_date, to_number
from excel_agent.schemas.workbook import (
    ApplyRequest,
    ApplyResponse,
    PreviewResponse,
    VersionInfo,
    WorkbookInfo,
)


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y") if value.time() == datetime.min.time() else value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if value is None or isinstance(value, bool | int | float | str):
        return value
    return str(value)


class WorkbookService:
    def __init__(self, root: Path, max_preview_rows: int = 500, max_apply_rows: int = 200) -> None:
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)
        self.max_preview_rows = max_preview_rows
        self.max_apply_rows = max_apply_rows
        self._locks: dict[str, threading.Lock] = {}
        self._locks_guard = threading.Lock()

    # ------------------------------------------------------------------ lưu trữ
    def _dir(self, workbook_id: str) -> Path:
        if not workbook_id.isalnum():
            raise NotFoundError(f"Mã workbook không hợp lệ: {workbook_id}")
        path = self.root / workbook_id
        if not path.is_dir():
            raise NotFoundError(f"Không tìm thấy workbook {workbook_id}")
        return path

    def _lock(self, workbook_id: str) -> threading.Lock:
        with self._locks_guard:
            return self._locks.setdefault(workbook_id, threading.Lock())

    def _meta(self, workbook_id: str) -> dict[str, Any]:
        return json.loads((self._dir(workbook_id) / "meta.json").read_text(encoding="utf-8"))

    def _save_meta(self, workbook_id: str, meta: dict[str, Any]) -> None:
        (self._dir(workbook_id) / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), "utf-8")

    def file_path(self, workbook_id: str) -> Path:
        meta = self._meta(workbook_id)
        return self._dir(workbook_id) / f"v{meta['versions'][-1]['version']}.xlsx"

    def download_name(self, workbook_id: str) -> str:
        return self._meta(workbook_id)["name"]

    def create(self, content: bytes, name: str, source: dict[str, Any] | None = None) -> WorkbookInfo:
        workbook_id = uuid.uuid4().hex[:12]
        folder = self.root / workbook_id
        folder.mkdir(parents=True)
        (folder / "v1.xlsx").write_bytes(content)
        try:
            load_workbook(folder / "v1.xlsx").close()
        except Exception as exc:  # file hỏng → không nhận
            raise ValidationError(f"File Excel không hợp lệ: {exc}") from exc
        meta = {
            "id": workbook_id,
            "name": name,
            "created_at": _now(),
            "source": source or {},
            "versions": [{"version": 1, "created_at": _now(), "action": "create", "summary": "Tạo file"}],
        }
        self._save_meta(workbook_id, meta)
        return self.info(workbook_id)

    def list(self, limit: int = 20) -> list[WorkbookInfo]:
        folders = sorted(
            (p for p in self.root.iterdir() if (p / "meta.json").exists()), key=lambda p: p.stat().st_mtime
        )
        return [self.info(p.name) for p in reversed(folders[-limit:])]

    # ------------------------------------------------------------------ đọc
    def info(self, workbook_id: str) -> WorkbookInfo:
        meta = self._meta(workbook_id)
        wb = load_workbook(self.file_path(workbook_id))
        try:
            ws = wb.worksheets[0]
            columns = _headers(ws)
            rows = max(ws.max_row - 1, 0)
            sheets = wb.sheetnames
        finally:
            wb.close()
        return WorkbookInfo(
            id=workbook_id,
            name=meta["name"],
            version=meta["versions"][-1]["version"],
            sheets=sheets,
            rows=rows,
            columns=columns,
            created_at=meta["created_at"],
            source=meta.get("source", {}),
            versions=[VersionInfo(**v) for v in meta["versions"]],
            preview_url=f"/api/v1/workbooks/{workbook_id}/preview",
            download_url=f"/api/v1/workbooks/{workbook_id}/download",
        )

    def preview(self, workbook_id: str, sheet: str | None = None, offset: int = 0, limit: int = 200) -> PreviewResponse:
        meta = self._meta(workbook_id)
        limit = max(1, min(limit, self.max_preview_rows))
        wb = load_workbook(self.file_path(workbook_id))
        try:
            ws = _sheet(wb, sheet)
            columns = _headers(ws)
            total = max(ws.max_row - 1, 0)
            rows = [
                [_json_value(v) for v in row]
                for row in ws.iter_rows(
                    min_row=2 + offset,
                    max_row=min(ws.max_row, 1 + offset + limit),
                    max_col=len(columns),
                    values_only=True,
                )
            ]
            return PreviewResponse(
                id=workbook_id,
                name=meta["name"],
                version=meta["versions"][-1]["version"],
                sheet=ws.title,
                sheets=wb.sheetnames,
                columns=columns,
                rows=rows,
                total_rows=total,
                offset=offset,
                limit=limit,
            )
        finally:
            wb.close()

    # ------------------------------------------------------------------ ghi
    def apply(self, workbook_id: str, req: ApplyRequest) -> ApplyResponse:
        with self._lock(workbook_id):
            meta = self._meta(workbook_id)
            current = meta["versions"][-1]["version"]
            if req.expected_version is not None and req.expected_version != current:
                raise ValidationError(f"File đã đổi sang phiên bản {current}; tải lại dữ liệu rồi áp lại")
            wb = load_workbook(self.file_path(workbook_id))
            try:
                ws = _sheet(wb, req.sheet)
                if req.mode == "set_column":
                    changed, column, message, warnings = self._set_column(ws, req)
                else:
                    changed, column, message, warnings = self._append_rows(ws, req)
                new_version = current + 1
                wb.save(self._dir(workbook_id) / f"v{new_version}.xlsx")
            finally:
                wb.close()
            meta["versions"].append(
                {"version": new_version, "created_at": _now(), "action": req.mode, "summary": message}
            )
            self._save_meta(workbook_id, meta)
        return ApplyResponse(
            workbook=self.info(workbook_id),
            changed_rows=changed,
            changed_column=column,
            message=message,
            warnings=warnings,
        )

    def undo(self, workbook_id: str) -> WorkbookInfo:
        with self._lock(workbook_id):
            meta = self._meta(workbook_id)
            if len(meta["versions"]) < 2:
                raise ValidationError("Không còn thao tác nào để hoàn tác")
            last = meta["versions"].pop()
            (self._dir(workbook_id) / f"v{last['version']}.xlsx").unlink(missing_ok=True)
            self._save_meta(workbook_id, meta)
        return self.info(workbook_id)

    def _set_column(self, ws: Worksheet, req: ApplyRequest) -> tuple[list[int], str, str, list[str]]:
        headers = _headers(ws)
        target = find_name(headers, req.target_column)
        warnings: list[str] = []
        if target is None:
            if not (req.create_column and req.target_column and req.target_column.strip()):
                raise ValidationError(f"Bảng không có cột '{req.target_column}'")
            target = req.target_column.strip()
            ws.cell(row=1, column=len(headers) + 1, value=target)
            headers.append(target)
            warnings.append(f"Đã tạo cột mới '{target}'")
        col_idx = headers.index(target) + 1
        field = find_name(list(req.records[0].keys()), req.source_field or target)
        if field is None:
            raise ValidationError(f"Kết quả tìm kiếm không có trường '{req.source_field or target}'")
        rows = sorted(set(req.row_indexes))
        if not rows:
            raise ValidationError("Chưa chọn dòng nào trong bảng Excel để ghi (bấm số thứ tự dòng để chọn)")
        if len(rows) > self.max_apply_rows:
            raise ValidationError(f"Mỗi lần ghi tối đa {self.max_apply_rows} dòng")
        n_data = max(ws.max_row - 1, 0)
        bad = [i for i in rows if i < 0 or i >= n_data]
        if bad:
            raise ValidationError(f"Dòng ngoài phạm vi dữ liệu: {bad}")
        if len(req.records) == 1:
            values = [req.records[0].get(field)] * len(rows)
        elif len(req.records) == len(rows):
            values = [r.get(field) for r in req.records]
        else:
            raise ValidationError(
                f"Đã chọn {len(req.records)} kết quả cho {len(rows)} dòng: chọn 1 kết quả (ghi cùng giá trị) "
                "hoặc số kết quả bằng số dòng (ghi lần lượt)"
            )
        kind = _column_kind(ws, col_idx)
        for row_idx, value in zip(rows, values, strict=True):
            cell = ws.cell(row=row_idx + 2, column=col_idx)
            cell.value = _coerce(value, kind, target)
            if kind == "date":
                cell.number_format = "dd/mm/yyyy"
        shown = values[0] if len(set(map(str, values))) == 1 else f"{len(values)} giá trị"
        message = f"Ghi '{shown}' (trường {field}) vào cột {target} cho {len(rows)} dòng"
        return rows, target, message, warnings

    def _append_rows(self, ws: Worksheet, req: ApplyRequest) -> tuple[list[int], str | None, str, list[str]]:
        headers = _headers(ws)
        if not headers:
            raise ValidationError("Sheet chưa có dòng tiêu đề")
        if len(req.records) > self.max_apply_rows:
            raise ValidationError(f"Mỗi lần thêm tối đa {self.max_apply_rows} dòng")
        record_keys = list(req.records[0].keys())
        mapping: dict[str, str] = {}
        for header in headers:
            wanted = (req.mapping or {}).get(header, header)
            field = find_name(record_keys, wanted)
            if field:
                mapping[header] = field
        if not mapping:
            raise ValidationError("Không trường nào của kết quả tìm kiếm khớp tên cột trong bảng")
        warnings = []
        unused = [k for k in record_keys if k not in mapping.values()]
        if unused:
            warnings.append(f"Bỏ qua trường không có cột tương ứng: {', '.join(unused)}")
        missing = [h for h in headers if h not in mapping]
        if missing:
            warnings.append(f"Cột để trống vì kết quả không có dữ liệu: {', '.join(missing)}")

        # Chống trùng khóa: cột đầu tiên là cột mã (ví dụ Mã KH) thì không thêm mã đã có
        key_header = headers[0]
        if key_header in mapping and normalize(key_header).startswith("ma"):
            existing = {normalize(c.value) for c in next(ws.iter_cols(min_col=1, max_col=1, min_row=2), ())} - {""}
            dup = [
                str(r.get(mapping[key_header]))
                for r in req.records
                if normalize(r.get(mapping[key_header])) in existing
            ]
            if dup:
                raise ValidationError(f"{key_header} đã có trong bảng: {', '.join(dup)}")

        kinds = {h: _column_kind(ws, i + 1) for i, h in enumerate(headers)}
        old_last = ws.max_row
        for record in req.records:
            values = [_coerce(record.get(mapping[h]), kinds[h], h) if h in mapping else None for h in headers]
            ws.append(values)
            new_row = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                above = ws.cell(row=new_row - 1, column=col_idx)
                if new_row - 1 >= 2:
                    ws.cell(row=new_row, column=col_idx).number_format = above.number_format
        _extend_tables(ws, old_last, ws.max_row)
        changed = list(range(old_last - 1, ws.max_row - 1))
        return changed, None, f"Thêm {len(req.records)} dòng mới (khớp {len(mapping)}/{len(headers)} cột)", warnings


# ---------------------------------------------------------------------- tiện ích
def _sheet(wb: Workbook, sheet: str | None) -> Worksheet:
    if sheet is None:
        return wb.worksheets[0]
    name = find_name(wb.sheetnames, sheet)
    if name is None:
        raise NotFoundError(f"Không có sheet '{sheet}'")
    return wb[name]


def _headers(ws: Worksheet) -> list[str]:
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]] if ws.max_row >= 1 else []
    while headers and headers[-1] == "":
        headers.pop()
    return headers


def _column_kind(ws: Worksheet, col_idx: int) -> str:
    if ws.max_row < 2:
        return "text"
    # Không đọc quá ws.max_row: openpyxl tự tạo ô rỗng khi duyệt ngoài vùng, làm lệch dòng cuối
    last = min(ws.max_row, 51)
    values = [c.value for c in next(ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=last), ())]
    values = [v for v in values if v is not None]
    if any(isinstance(v, datetime | date) for v in values):
        return "date"
    if values and all(isinstance(v, int | float) and not isinstance(v, bool) for v in values):
        return "number"
    return "text"


def _coerce(value: Any, kind: str, column: str) -> Any:
    if value is None or value == "":
        return None
    if kind == "date":
        parsed = parse_date(value)
        if parsed is None:
            raise ValidationError(f"'{value}' không phải ngày hợp lệ cho cột {column}")
        return parsed
    if kind == "number":
        number = to_number(value)
        if number is None:
            raise ValidationError(f"'{value}' không phải số cho cột {column}")
        return int(number) if number.is_integer() else number
    return value


def _extend_tables(ws: Worksheet, old_last: int, new_last: int) -> None:
    """Giữ các dòng mới nằm trong Excel Table (ListObject) nếu bảng kết thúc ở dòng cuối cũ."""
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if max_row == old_last:
            table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{new_last}"
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
