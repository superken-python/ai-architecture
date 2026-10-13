"""Logic của API dữ liệu giả lập: đọc dataset JSON, lọc, tìm kiếm, xuất file Excel."""

from __future__ import annotations

import io
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_agent.core.exceptions import NotFoundError, ValidationError
from excel_agent.core.text import find_name, normalize
from excel_agent.core.values import parse_date, to_number
from excel_agent.schemas.data_api import DatasetInfo, ExportRequest, Filter, SearchResponse

DATA_DIR = Path(__file__).parent / "data"
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


class DataRepository:
    """Kho dữ liệu chỉ đọc, nạp một lần từ thư mục `data/`."""

    def __init__(self, data_dir: Path = DATA_DIR) -> None:
        catalog = yaml.safe_load((data_dir / "datasets.yaml").read_text(encoding="utf-8"))
        self._infos: dict[str, DatasetInfo] = {}
        self._rows: dict[str, list[dict[str, Any]]] = {}
        for item in catalog["datasets"]:
            rows = json.loads((data_dir / f"{item['name']}.json").read_text(encoding="utf-8"))
            self._infos[item["name"]] = DatasetInfo(**item, rows=len(rows))
            self._rows[item["name"]] = rows

    # ---------- danh mục ----------
    def list(self) -> list[DatasetInfo]:
        return list(self._infos.values())

    def info(self, dataset: str) -> DatasetInfo:
        key = normalize(dataset).replace(" ", "_")
        for info in self._infos.values():
            if key in (info.name, normalize(info.title).replace(" ", "_")):
                return info
        raise NotFoundError(f"Không có dataset '{dataset}'. Có: {', '.join(self._infos)}")

    def _column(self, info: DatasetInfo, column: str) -> str:
        name = find_name([c.name for c in info.columns], column)
        if name is None:
            raise NotFoundError(f"Dataset '{info.name}' không có cột '{column}'")
        return name

    # ---------- lọc ----------
    def _match(self, info: DatasetInfo, row: dict[str, Any], flt: Filter) -> bool:
        column = self._column(info, flt.column)
        ctype = next(c.type for c in info.columns if c.name == column)
        cell = row.get(column)
        if flt.op == "eq":
            return normalize(cell) == normalize(flt.value)
        if flt.op == "contains":
            return normalize(flt.value) in normalize(cell)
        if ctype == "date":
            left, right = parse_date(cell), parse_date(flt.value)
        else:
            left, right = to_number(cell), to_number(flt.value)
        if left is None or right is None:
            raise ValidationError(f"Không so sánh '{flt.op}' được giữa '{cell}' và '{flt.value}' (cột {column})")
        return {"gt": left > right, "gte": left >= right, "lt": left < right, "lte": left <= right}[flt.op]

    def query(self, dataset: str, filters: list[Filter] | None = None, limit: int = 1000) -> list[dict[str, Any]]:
        info = self.info(dataset)
        rows = [r for r in self._rows[info.name] if all(self._match(info, r, f) for f in filters or [])]
        return rows[:limit]

    # ---------- tìm kiếm ----------
    def search(self, dataset: str, query: str, column: str | None = None, limit: int = 20) -> SearchResponse:
        info = self.info(dataset)
        col = self._column(info, column) if column else None
        tokens = normalize(query).split()
        scored: list[tuple[int, dict[str, Any]]] = []
        for row in self._rows[info.name]:
            haystack = normalize(row.get(col)) if col else " | ".join(normalize(v) for v in row.values())
            if not all(t in haystack for t in tokens):
                continue
            exact = col is not None and haystack == normalize(query)
            scored.append((0 if exact else 1, row))
        scored.sort(key=lambda item: item[0])
        records = [row for _, row in scored[:limit]]
        return SearchResponse(
            dataset=info.name,
            query=query,
            column=col,
            total=len(scored),
            columns=[c.name for c in info.columns],
            records=records,
        )

    # ---------- xuất Excel ----------
    def export_xlsx(self, req: ExportRequest) -> tuple[bytes, str, int]:
        info = self.info(req.dataset)
        rows = self.query(info.name, req.filters, req.limit)
        columns = [self._column(info, c) for c in req.columns] if req.columns else [c.name for c in info.columns]
        types = {c.name: c.type for c in info.columns}
        title = req.title or info.title

        wb = Workbook()
        ws = wb.active
        ws.title = _SHEET_FORBIDDEN.sub(" ", title)[:31].strip() or info.title
        ws.append(columns)
        for row in rows:
            ws.append([_excel_value(row.get(c), types[c]) for c in columns])
        _format_sheet(ws, columns, types, len(rows), table_name=f"tbl_{info.name}")

        meta = wb.create_sheet("Thông tin")
        filters = "; ".join(f"{f.column} {f.op} {f.value}" for f in req.filters) or "(không)"
        for key, value in [
            ("Nguồn", "API dữ liệu giả lập (POC CMB-10)"),
            ("Dataset", f"{info.title} ({info.name})"),
            ("Bộ lọc", filters),
            ("Số dòng", len(rows)),
            ("Thời điểm tạo", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ]:
            meta.append([key, value])
        meta.column_dimensions["A"].width = 16
        meta.column_dimensions["B"].width = 60

        buffer = io.BytesIO()
        wb.save(buffer)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return buffer.getvalue(), f"{info.name}_{stamp}.xlsx", len(rows)


def _excel_value(value: Any, ctype: str) -> Any:
    if value is None:
        return None
    if ctype == "date":
        return parse_date(value) or value
    return value


def _format_sheet(ws: Any, columns: list[str], types: dict[str, str], n_rows: int, table_name: str) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="217346")
    for idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font, cell.fill = header_font, header_fill
        fmt = {"money": "#,##0", "number": "#,##0", "date": "dd/mm/yyyy"}.get(types.get(name, "text"))
        if fmt:
            for row in range(2, n_rows + 2):
                ws.cell(row=row, column=idx).number_format = fmt
        width = max(
            [len(name)] + [len(str(ws.cell(row=r, column=idx).value or "")) for r in range(2, min(n_rows, 200) + 2)]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, width + 2), 45)
    ws.freeze_panes = "A2"
    if n_rows > 0:  # Excel Table cần ít nhất một dòng dữ liệu
        ref = f"A1:{get_column_letter(len(columns))}{n_rows + 1}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
